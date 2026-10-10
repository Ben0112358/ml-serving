import io
import zipfile
from dataclasses import dataclass
from html import escape
from xml.etree import ElementTree

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt
from odf import teletype
from odf.opendocument import OpenDocumentText, load
from odf.style import PageLayout, PageLayoutProperties, Style, TextProperties
from odf.text import H, P
from pypdf import PdfReader
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

ALLOWED_SUFFIXES = {".docx", ".pdf", ".odt"}
_DOCX_MEDIA = (
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
)
_ODT_MEDIA = "application/vnd.oasis.opendocument.text"
_THEME_NS = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}


@dataclass
class DocumentStyle:
    font_name: str | None = None
    heading_font_name: str | None = None
    font_size_pt: float | None = None
    heading_size_pt: float | None = None
    page_width_pt: float | None = None
    page_height_pt: float | None = None
    margin_top_pt: float | None = None
    margin_bottom_pt: float | None = None
    margin_left_pt: float | None = None
    margin_right_pt: float | None = None


def suffix_of(filename: str) -> str:
    lower = filename.lower()
    for suffix in ALLOWED_SUFFIXES:
        if lower.endswith(suffix):
            return suffix
    raise ValueError("Use a .docx, .pdf, or .odt file")


def extract_text(filename: str, data: bytes) -> str:
    suffix = suffix_of(filename)
    if suffix == ".docx":
        document = Document(io.BytesIO(data))
        return "\n".join(paragraph.text for paragraph in document.paragraphs)
    if suffix == ".pdf":
        reader = PdfReader(io.BytesIO(data))
        return "\n".join(page.extract_text() or "" for page in reader.pages)
    document = load(io.BytesIO(data))
    paragraphs = document.getElementsByType(P)
    return "\n".join(teletype.extractText(item) for item in paragraphs)


def extract_style(filename: str, data: bytes) -> DocumentStyle:
    suffix = suffix_of(filename)
    if suffix == ".docx":
        return _docx_style(data)
    if suffix == ".pdf":
        return _pdf_style(data)
    return _odt_style(data)


def build_document(
    filename: str,
    body: str,
    template_name: str | None = None,
    template: bytes | None = None,
) -> tuple[bytes, str]:
    suffix = suffix_of(filename)
    if template and template_name and suffix_of(template_name) == suffix:
        if suffix == ".docx":
            return _clone_docx(template, body), _DOCX_MEDIA
        if suffix == ".odt":
            return _clone_odt(template, body), _ODT_MEDIA
    profile = DocumentStyle()
    if template and template_name:
        profile = extract_style(template_name, template)
    if suffix == ".docx":
        return _fresh_docx(body, profile), _DOCX_MEDIA
    if suffix == ".pdf":
        return _fresh_pdf(body, profile), "application/pdf"
    return _fresh_odt(body, profile), _ODT_MEDIA


def _heading_level(line: str) -> int | None:
    if line.startswith("### "):
        return 3
    if line.startswith("## "):
        return 2
    if line.startswith("# "):
        return 1
    return None


def _heading_text(line: str, level: int) -> str:
    marker = f"{'#' * level} "
    return line.removeprefix(marker).strip()


def _clone_docx(template: bytes, body: str) -> bytes:
    document = Document(io.BytesIO(template))
    body_element = document.element.body
    for child in list(body_element):
        if child.tag != qn("w:sectPr"):
            body_element.remove(child)
    for section in document.sections:
        for tag in ("headerReference", "footerReference"):
            for node in section._sectPr.findall(qn(f"w:{tag}")):
                section._sectPr.remove(node)
    for line in body.splitlines():
        _add_docx_line(document, line, DocumentStyle())
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _fresh_docx(body: str, profile: DocumentStyle) -> bytes:
    document = Document()
    _apply_docx_page(document, profile)
    for line in body.splitlines():
        _add_docx_line(document, line, profile)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _apply_docx_page(document: Document, profile: DocumentStyle) -> None:
    section = document.sections[0]
    if profile.page_width_pt:
        section.page_width = Pt(profile.page_width_pt)
    if profile.page_height_pt:
        section.page_height = Pt(profile.page_height_pt)
    if profile.margin_top_pt:
        section.top_margin = Pt(profile.margin_top_pt)
    if profile.margin_bottom_pt:
        section.bottom_margin = Pt(profile.margin_bottom_pt)
    if profile.margin_left_pt:
        section.left_margin = Pt(profile.margin_left_pt)
    if profile.margin_right_pt:
        section.right_margin = Pt(profile.margin_right_pt)


def _add_docx_line(
    document: Document,
    line: str,
    profile: DocumentStyle,
) -> None:
    level = _heading_level(line)
    text = _heading_text(line, level) if level else line
    style_name = f"Heading {level}" if level else None
    if style_name and _has_style(document, style_name):
        paragraph = document.add_paragraph(text, style=style_name)
    else:
        paragraph = document.add_paragraph(text)
        if level:
            for run in paragraph.runs:
                run.bold = True
    font_name = profile.heading_font_name if level else profile.font_name
    size = profile.heading_size_pt if level else profile.font_size_pt
    if not font_name and not size:
        return
    for run in paragraph.runs:
        if font_name:
            run.font.name = font_name
        if size:
            run.font.size = Pt(size)


def _has_style(document: Document, name: str) -> bool:
    try:
        document.styles[name]
    except KeyError:
        return False
    return True


def _clone_odt(template: bytes, body: str) -> bytes:
    document = load(io.BytesIO(template))
    body_style = _common_odt_style(document)
    heading_style = _odt_heading_style(document)
    text = document.text
    for child in list(text.childNodes):
        text.removeChild(child)
    for line in body.splitlines():
        level = _heading_level(line)
        content = _heading_text(line, level) if level else line
        if level and heading_style:
            text.addElement(
                H(outlinelevel=level, stylename=heading_style, text=content)
            )
            continue
        paragraph = P(text=content)
        if body_style:
            paragraph.setAttribute("stylename", body_style)
        text.addElement(paragraph)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _fresh_odt(body: str, profile: DocumentStyle) -> bytes:
    document = OpenDocumentText()
    if profile.font_name or profile.font_size_pt:
        style = Style(name="BoardBody", family="paragraph")
        properties = TextProperties()
        if profile.font_name:
            properties.setAttribute("fontname", profile.font_name)
        if profile.font_size_pt:
            properties.setAttribute("fontsize", f"{profile.font_size_pt}pt")
        style.addElement(properties)
        document.styles.addElement(style)
    for line in body.splitlines():
        level = _heading_level(line)
        content = _heading_text(line, level) if level else line
        paragraph = P(text=content)
        if profile.font_name or profile.font_size_pt:
            paragraph.setAttribute("stylename", "BoardBody")
        document.text.addElement(paragraph)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _fresh_pdf(body: str, profile: DocumentStyle) -> bytes:
    size = profile.font_size_pt or 11
    heading_size = profile.heading_size_pt or size + 2
    body_style = ParagraphStyle(
        "BoardBody",
        fontName=_reportlab_font(profile.font_name, bold=False),
        fontSize=size,
        leading=size * 1.3,
    )
    heading_style = ParagraphStyle(
        "BoardHeading",
        fontName=_reportlab_font(
            profile.heading_font_name or profile.font_name,
            bold=True,
        ),
        fontSize=heading_size,
        leading=heading_size * 1.2,
        spaceBefore=8,
        spaceAfter=4,
    )
    pagesize = A4
    if profile.page_width_pt and profile.page_height_pt:
        pagesize = (profile.page_width_pt, profile.page_height_pt)
    buffer = io.BytesIO()
    story = []
    for line in body.splitlines():
        level = _heading_level(line)
        content = escape(_heading_text(line, level) if level else line)
        story.append(
            Paragraph(
                content or "&nbsp;", heading_style if level else body_style
            )
        )
        story.append(Spacer(1, 6))
    SimpleDocTemplate(
        buffer,
        pagesize=pagesize,
        leftMargin=profile.margin_left_pt or 72,
        rightMargin=profile.margin_right_pt or 72,
        topMargin=profile.margin_top_pt or 72,
        bottomMargin=profile.margin_bottom_pt or 72,
    ).build(story)
    return buffer.getvalue()


def _reportlab_font(name: str | None, bold: bool) -> str:
    cleaned = (name or "").split("+")[-1].lower().replace("_", " ")
    sans = ("arial", "calibri", "helvetica", "carlito", "liberation sans")
    mono = ("courier", "consolas", "mono")
    if any(token in cleaned for token in sans):
        return "Helvetica-Bold" if bold else "Helvetica"
    if any(token in cleaned for token in mono):
        return "Courier-Bold" if bold else "Courier"
    return "Times-Bold" if bold else "Times-Roman"


def _docx_style(data: bytes) -> DocumentStyle:
    document = Document(io.BytesIO(data))
    section = document.sections[0]
    normal = document.styles["Normal"].font
    heading = (
        document.styles["Heading 1"].font
        if _has_style(document, "Heading 1")
        else None
    )
    run_name, run_size = _dominant_run_font(document)
    return DocumentStyle(
        font_name=normal.name or run_name or _theme_font(data, minor=True),
        heading_font_name=(
            (heading.name if heading else None)
            or _theme_font(data, minor=False)
        ),
        font_size_pt=_length_pt(normal.size) or run_size,
        heading_size_pt=_length_pt(heading.size) if heading else None,
        page_width_pt=_length_pt(section.page_width),
        page_height_pt=_length_pt(section.page_height),
        margin_top_pt=_length_pt(section.top_margin),
        margin_bottom_pt=_length_pt(section.bottom_margin),
        margin_left_pt=_length_pt(section.left_margin),
        margin_right_pt=_length_pt(section.right_margin),
    )


def _dominant_run_font(document: Document) -> tuple[str | None, float | None]:
    counts: dict[tuple[str | None, float | None], int] = {}
    for paragraph in document.paragraphs:
        for run in paragraph.runs:
            name = run.font.name
            size = _length_pt(run.font.size)
            if not name and size is None:
                continue
            key = (name, size)
            counts[key] = counts.get(key, 0) + max(len(run.text), 1)
    if not counts:
        return None, None
    name, size = max(counts, key=counts.get)
    return name, size


def _theme_font(data: bytes, minor: bool) -> str | None:
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        if "word/theme/theme1.xml" not in archive.namelist():
            return None
        root = ElementTree.fromstring(archive.read("word/theme/theme1.xml"))
    tag = "a:minorFont" if minor else "a:majorFont"
    node = root.find(f".//{tag}/a:latin", _THEME_NS)
    if node is None:
        return None
    return node.get("typeface")


def _pdf_style(data: bytes) -> DocumentStyle:
    reader = PdfReader(io.BytesIO(data))
    page = reader.pages[0]
    box = page.mediabox
    return DocumentStyle(
        font_name=_pdf_font_name(page),
        page_width_pt=float(box.width),
        page_height_pt=float(box.height),
    )


def _pdf_font_name(page) -> str | None:
    resources = page.get("/Resources") or {}
    if hasattr(resources, "get_object"):
        resources = resources.get_object()
    fonts = resources.get("/Font") if resources else None
    if fonts is None:
        return None
    if hasattr(fonts, "get_object"):
        fonts = fonts.get_object()
    for font in fonts.values():
        obj = font.get_object() if hasattr(font, "get_object") else font
        base = str(obj.get("/BaseFont", ""))
        if "+" in base:
            base = base.split("+", 1)[1]
        return base.replace("-", " ") or None
    return None


def _odt_style(data: bytes) -> DocumentStyle:
    document = load(io.BytesIO(data))
    profile = DocumentStyle(font_name=_common_odt_font(document))
    for layout in document.automaticstyles.getElementsByType(PageLayout):
        props = layout.getElementsByType(PageLayoutProperties)
        if not props:
            continue
        page = props[0]
        profile.page_width_pt = _css_pt(page.getAttribute("pagewidth"))
        profile.page_height_pt = _css_pt(page.getAttribute("pageheight"))
        profile.margin_top_pt = _css_pt(page.getAttribute("margintop"))
        profile.margin_bottom_pt = _css_pt(page.getAttribute("marginbottom"))
        profile.margin_left_pt = _css_pt(page.getAttribute("marginleft"))
        profile.margin_right_pt = _css_pt(page.getAttribute("marginright"))
        break
    return profile


def _common_odt_style(document) -> str | None:
    counts: dict[str, int] = {}
    for paragraph in document.getElementsByType(P):
        name = paragraph.getAttribute("stylename")
        if name:
            counts[str(name)] = counts.get(str(name), 0) + 1
    if not counts:
        return None
    return max(counts, key=counts.get)


def _common_odt_font(document) -> str | None:
    name = _common_odt_style(document)
    if not name:
        return None
    for style in document.styles.getElementsByType(Style):
        if style.getAttribute("name") != name:
            continue
        props = style.getElementsByType(TextProperties)
        if props:
            return props[0].getAttribute("fontname")
    return None


def _odt_heading_style(document) -> str | None:
    for style in document.styles.getElementsByType(Style):
        name = str(style.getAttribute("name") or "")
        family = style.getAttribute("family")
        if family == "paragraph" and name.lower().startswith("heading"):
            return name
    return None


def _length_pt(length) -> float | None:
    if length is None:
        return None
    return float(length.pt)


def _css_pt(value) -> float | None:
    if not value:
        return None
    text = str(value).strip().lower()
    number = ""
    for char in text:
        if char.isdigit() or char == ".":
            number += char
        elif number:
            break
    if not number:
        return None
    amount = float(number)
    if text.endswith("cm"):
        return amount * 72 / 2.54
    if text.endswith("mm"):
        return amount * 72 / 25.4
    if text.endswith("in"):
        return amount * 72
    return amount
