import io

from docx import Document
from docx.shared import Inches, Pt
from odf.opendocument import OpenDocumentText, load
from odf.style import Style, TextProperties
from odf.text import P
from pypdf import PdfReader

from ml_serving.board_meeting_protocol.documents import (
    build_document,
    extract_text,
)


def test_docx_roundtrip():
    raw, _media = build_document("protocol.docx", "Decision one\nDecision two")
    text = extract_text("protocol.docx", raw)
    assert "Decision one" in text
    assert "Decision two" in text


def test_docx_keeps_finished_font_and_margins():
    template = Document()
    template.styles["Normal"].font.name = "Calibri"
    template.styles["Normal"].font.size = Pt(14)
    section = template.sections[0]
    section.top_margin = Inches(0.6)
    section.left_margin = Inches(0.8)
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    template.add_paragraph("Old protocol text")
    buffer = io.BytesIO()
    template.save(buffer)
    raw, _media = build_document(
        "new.docx",
        "# Budget vote\nThe board approved it.",
        "finished.docx",
        buffer.getvalue(),
    )
    result = Document(io.BytesIO(raw))
    assert result.paragraphs[0].text == "Budget vote"
    assert result.paragraphs[1].text == "The board approved it."
    assert "Old protocol text" not in extract_text("new.docx", raw)
    assert result.sections[0].top_margin == section.top_margin
    assert result.sections[0].left_margin == section.left_margin
    assert result.styles["Normal"].font.size == Pt(14)


def test_pdf_uses_finished_page_size_and_font():
    template = Document()
    template.styles["Normal"].font.name = "Calibri"
    section = template.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    template.add_paragraph("Finished")
    buffer = io.BytesIO()
    template.save(buffer)
    pdf, _media = build_document(
        "notes.pdf",
        "A short note",
        "finished.docx",
        buffer.getvalue(),
    )
    page = PdfReader(io.BytesIO(pdf)).pages[0]
    assert abs(float(page.mediabox.width) - 8.5 * 72) < 2
    assert b"Helvetica" in pdf


def test_odt_keeps_finished_paragraph_style():
    template = OpenDocumentText()
    style = Style(name="BoardBody", family="paragraph")
    properties = TextProperties()
    properties.setAttribute("fontname", "Calibri")
    properties.setAttribute("fontsize", "14pt")
    style.addElement(properties)
    template.styles.addElement(style)
    template.text.addElement(P(stylename="BoardBody", text="Old protocol"))
    buffer = io.BytesIO()
    template.save(buffer)
    raw, _media = build_document(
        "new.odt",
        "The board approved it.",
        "finished.odt",
        buffer.getvalue(),
    )
    document = load(io.BytesIO(raw))
    paragraphs = document.getElementsByType(P)
    assert len(paragraphs) == 1
    assert paragraphs[0].getAttribute("stylename") == "BoardBody"
    assert "Old protocol" not in extract_text("new.odt", raw)
    assert "approved" in extract_text("new.odt", raw)


def test_pdf_and_odt_build():
    pdf, pdf_type = build_document("notes.pdf", "A short note")
    odt, odt_type = build_document("notes.odt", "A short note")
    assert pdf.startswith(b"%PDF")
    assert pdf_type == "application/pdf"
    assert odt[:2] == b"PK"
    assert "opendocument" in odt_type
